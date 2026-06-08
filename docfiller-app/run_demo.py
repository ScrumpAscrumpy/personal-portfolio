#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Automated demo: launch DocFiller, drive it, take screenshots."""
import os, sys, time, subprocess
from pathlib import Path

DEMO_DIR = Path("/tmp/docfiller_demo")
DEMO_DIR.mkdir(exist_ok=True)
SHOT_DIR = Path("/home/user/personal-portfolio/docfiller-app/demo_screenshots")
SHOT_DIR.mkdir(exist_ok=True)

# ── 1. Create a sample Word form ─────────────────────────────────────────────
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def make_sample():
    path = str(DEMO_DIR / "退休申请表.docx")
    doc = Document()
    doc.add_heading("退休金领取资格认证申请表", level=1).alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=7, cols=4)
    table.style = "Table Grid"
    rows_data = [
        ("姓名",     "",              "性别",     ""),
        ("出生日期", "年    月    日", "民族",     ""),
        ("身份证号", "",              "联系电话", ""),
        ("户籍地址", "",              "",          ""),
        ("工作单位", "",              "职务",      ""),
        ("退休日期", "年    月    日", "社保编号", ""),
        ("备注",     "",              "",           ""),
    ]
    for i, (a, b, c, d) in enumerate(rows_data):
        r = table.rows[i]
        r.cells[0].text = a
        r.cells[1].text = b
        r.cells[2].text = c
        r.cells[3].text = d

    doc.save(path)
    return path

doc_path = make_sample()
print(f"Sample doc: {doc_path}")

# ── 2. Boot the app in the same process ──────────────────────────────────────
sys.path.insert(0, str(Path(__file__).parent))

import importlib.util
spec = importlib.util.spec_from_file_location(
    "docfiller", Path(__file__).parent / "docfiller.py")
mod = importlib.util.module_from_spec(spec)
spec.loader.exec_module(mod)

DocFillerApp   = mod.DocFillerApp
analyze_document = mod.analyze_document
fill_document    = mod.fill_document

import tkinter as tk

shot_n = [0]

def screenshot(name):
    shot_n[0] += 1
    out = str(SHOT_DIR / f"{shot_n[0]:02d}_{name}.png")
    subprocess.run(["scrot", "-d", "0", out], env={**os.environ, "DISPLAY": ":99"})
    print(f"  📸  {out}")
    return out

# ── 3. Build app & inject automation ─────────────────────────────────────────
app = DocFillerApp.__new__(DocFillerApp)
import tkinter as tk
from tkinter import ttk
app.root = tk.Tk()
app.root.title("DocFiller · Word表格智能填写")
app.root.geometry("980x740")
app.root.minsize(800, 580)
app.root.configure(bg=mod.C["bg"])
app.cfg      = mod.load_config()
app.doc_path = None
app.fields   = []
app._vars    = {}
mod.apply_styles()
app._build_ui = mod.DocFillerApp._build_ui.__get__(app, mod.DocFillerApp)
app._build_fill_tab    = mod.DocFillerApp._build_fill_tab.__get__(app, mod.DocFillerApp)
app._build_history_tab = mod.DocFillerApp._build_history_tab.__get__(app, mod.DocFillerApp)
app._open_doc          = mod.DocFillerApp._open_doc.__get__(app, mod.DocFillerApp)
app._render_form       = mod.DocFillerApp._render_form.__get__(app, mod.DocFillerApp)
app._generate          = mod.DocFillerApp._generate.__get__(app, mod.DocFillerApp)
app._on_tab            = mod.DocFillerApp._on_tab.__get__(app, mod.DocFillerApp)
app._refresh_history   = mod.DocFillerApp._refresh_history.__get__(app, mod.DocFillerApp)
app._selected_history  = mod.DocFillerApp._selected_history.__get__(app, mod.DocFillerApp)
app._open_hist_folder  = mod.DocFillerApp._open_hist_folder.__get__(app, mod.DocFillerApp)
app._refill_from_history = mod.DocFillerApp._refill_from_history.__get__(app, mod.DocFillerApp)
app._open_settings     = mod.DocFillerApp._open_settings.__get__(app, mod.DocFillerApp)
app._pick_out_dir      = mod.DocFillerApp._pick_out_dir.__get__(app, mod.DocFillerApp)
app._clear_form        = mod.DocFillerApp._clear_form.__get__(app, mod.DocFillerApp)
app._open_path         = mod.DocFillerApp._open_path.__get__(app, mod.DocFillerApp)
app.run                = mod.DocFillerApp.run.__get__(app, mod.DocFillerApp)
app._build_ui()

SAMPLE_VALUES = {
    "姓名":     "张三",
    "性别":     "男",
    "出生日期": "1958年03月15日",
    "民族":     "汉族",
    "身份证号": "110101195803151234",
    "联系电话": "13800138000",
    "户籍地址": "北京市朝阳区建国路88号",
    "工作单位": "北京市第一机械厂",
    "职务":     "高级工程师",
    "退休日期": "2018年03月15日",
    "社保编号": "1101011234567",
    "备注":     "无",
}

def step_initial():
    """Step 1: Show initial empty state."""
    app.root.update()
    screenshot("01_初始界面")
    app.root.after(300, step_open_doc)

def step_open_doc():
    """Step 2: Load the sample document (bypass file dialog)."""
    app.doc_path = doc_path
    app._path_var.set(os.path.basename(doc_path))
    app._status_var.set(f"正在分析文档：{os.path.basename(doc_path)} ...")
    app.root.update()
    app.fields = analyze_document(doc_path)
    app._render_form(app.fields)
    app._gen_btn.config(text=f"生成文档  ({len(app.fields)} 个字段)")
    app._out_var.set(str(DEMO_DIR))
    app._status_var.set(f"已识别 {len(app.fields)} 个字段  ·  {os.path.basename(doc_path)}")
    app.root.update()
    screenshot("02_文档已加载_字段识别")
    app.root.after(300, step_fill_fields)

def step_fill_fields():
    """Step 3: Type sample values into every field."""
    for label, value in SAMPLE_VALUES.items():
        if label in app._vars:
            app._vars[label].set(value)
    app.root.update()
    screenshot("03_填写信息")
    app.root.after(300, step_generate)

def step_generate():
    """Step 4: Click Generate."""
    app._generate()
    app.root.update()
    screenshot("04_生成完成")
    app.root.after(500, step_history)

def step_history():
    """Step 5: Switch to History tab."""
    app.nb.select(1)
    app._refresh_history()
    app.root.update()
    screenshot("05_历史记录")
    app.root.after(300, done)

def done():
    print("\n✅ Demo complete. Screenshots saved to:", SHOT_DIR)
    app.root.destroy()

app.root.after(400, step_initial)
app.root.mainloop()
