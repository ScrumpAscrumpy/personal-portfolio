import os
import re
from typing import Dict, List
from docx import Document


def _replace_in_paragraph(para, replacements: Dict[str, str]):
    """Replace {{key}} placeholders in a paragraph, handling split runs."""
    full_text = para.text
    if not any(f'{{{{{k}}}}}' in full_text for k in replacements):
        return

    # Try simple per-run replacement first
    for run in para.runs:
        for key, value in replacements.items():
            ph = f'{{{{{key}}}}}'
            if ph in run.text:
                run.text = run.text.replace(ph, str(value) if value is not None else '')

    # If any placeholder is still split across runs, consolidate
    remaining = para.text
    if any(f'{{{{{k}}}}}' in remaining for k in replacements):
        new_text = remaining
        for key, value in replacements.items():
            new_text = new_text.replace(f'{{{{{key}}}}}', str(value) if value is not None else '')
        if para.runs:
            # Preserve font of first run, clear the rest
            para.runs[0].text = new_text
            for run in para.runs[1:]:
                run.text = ''


def _process_doc(doc: Document, replacements: Dict[str, str]):
    for para in doc.paragraphs:
        _replace_in_paragraph(para, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    _replace_in_paragraph(para, replacements)
                # Handle nested tables
                for nested in cell.tables:
                    for nrow in nested.rows:
                        for ncell in nrow.cells:
                            for para in ncell.paragraphs:
                                _replace_in_paragraph(para, replacements)

    for section in doc.sections:
        for para in section.header.paragraphs:
            _replace_in_paragraph(para, replacements)
        for para in section.footer.paragraphs:
            _replace_in_paragraph(para, replacements)
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        _replace_in_paragraph(para, replacements)
        for table in section.footer.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        _replace_in_paragraph(para, replacements)


def fill_document(template_path: str, replacements: Dict[str, str], output_path: str):
    """Fill a Word template with replacements and save to output_path."""
    doc = Document(template_path)
    _process_doc(doc, replacements)
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)


def extract_placeholders(docx_path: str) -> List[str]:
    """Return unique {{placeholder}} keys found in the Word document."""
    doc = Document(docx_path)
    texts = []

    for para in doc.paragraphs:
        texts.append(para.text)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for para in cell.paragraphs:
                    texts.append(para.text)
    for section in doc.sections:
        for para in section.header.paragraphs:
            texts.append(para.text)
        for para in section.footer.paragraphs:
            texts.append(para.text)

    full = '\n'.join(texts)
    found = re.findall(r'\{\{(.+?)\}\}', full)
    seen, unique = set(), []
    for k in found:
        k = k.strip()
        if k and k not in seen:
            seen.add(k)
            unique.append(k)
    return unique
